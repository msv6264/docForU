from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from io import BytesIO


def create_ppt(data):

    def style_run(run, size, bold=False, color=None):
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color

    def style_para(p, align):
        p.alignment = align

    prs = Presentation()

    BLUE = RGBColor(0, 102, 204)
    DARK = RGBColor(60, 60, 60)
    BLACK = RGBColor(0, 0, 0)

    # ---------------- TITLE SLIDE ----------------
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    title = slide.shapes.title
    
    # Position the title box horizontally and vertically in the middle
    title.left = Inches(0.5)
    title.top = Inches(3.0)
    title.width = Inches(9.0)
    title.height = Inches(1.5)

    title.text = data["title"]

    p = title.text_frame.paragraphs[0]
    style_para(p, PP_ALIGN.CENTER)
    run = p.runs[0]
    style_run(run, 44, True, BLUE)

    # ---------------- CONTENT ----------------
    for section in data["sections"]:

        for i, page in enumerate(section["pages"]):

            slide = prs.slides.add_slide(prs.slide_layouts[1])

            title_box = slide.shapes.title
            content_box = slide.placeholders[1]
            
            # Position content box
            content_box.left = Inches(1.0)
            content_box.top = Inches(2.0)
            content_box.width = Inches(8.0)

            # -------- TITLE --------
            if i == 0:
                title_box.text = section["sectionName"]
                p = title_box.text_frame.paragraphs[0]
                style_para(p, PP_ALIGN.CENTER)
                run = p.runs[0]
                style_run(run, 34, True, BLUE)
            else:
                title_box.text = page["pageTitle"]
                p = title_box.text_frame.paragraphs[0]
                style_para(p, PP_ALIGN.CENTER)
                run = p.runs[0]
                style_run(run, 28, True, DARK)

            # -------- CONTENT BOX --------
            tf = content_box.text_frame
            tf.clear()

            if i == 0:
                # -------- PAGE TITLE (SIDE HEADING) --------
                p = tf.paragraphs[0]
                p.text = page["pageTitle"]
                p.level = 0
                
                # Remove bullet point
                pPr = p._element.get_or_add_pPr()
                pPr.append(OxmlElement('a:buNone'))
                
                style_para(p, PP_ALIGN.LEFT)

                run = p.runs[0]
                style_run(run, 24, True, DARK)

                # -------- BULLET CONTENT --------
                for j, bullet in enumerate(page["content"]):
                    p = tf.add_paragraph()
                    p.text = bullet
                    p.level = 0
                    style_para(p, PP_ALIGN.LEFT)

                    run = p.runs[0]
                    style_run(run, 20, False, BLACK)
            else:
                # -------- BULLET CONTENT ONLY --------
                for j, bullet in enumerate(page["content"]):
                    if j == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                        
                    p.text = bullet
                    p.level = 0
                    style_para(p, PP_ALIGN.LEFT)

                    run = p.runs[0]
                    style_run(run, 20, False, BLACK)

    ppt_stream = BytesIO()
    prs.save(ppt_stream)
    ppt_stream.seek(0)

    return ppt_stream