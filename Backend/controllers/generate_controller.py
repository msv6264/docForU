import asyncio
from io import BytesIO
from fastapi import HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from services.ppt_service import create_ppt
from services.docx_service import create_docx
from ai.cerebras import prompt_to_cerebras
from services.promptService import getFilledPrompt
from services.cleanDocData import prepDocData
from utils.response_parser import parse_ai_json

from docx import Document as DocxDocument
from pptx import Presentation

# Concurrency limit: Only 1 document generation at a time
# This replaces the Node.js queueService.js logic
generation_queue = asyncio.Semaphore(1)

async def handle_generate_doc(doc_data: dict):
    """
    Controller logic for orchestrating document generation with a concurrency queue.
    """
    async with generation_queue:
        try:
            # 1. Prepare the document data (formatting)
            final_doc_obj = await prepDocData(doc_data)
            
            # 2. Fill the prompt template with processed data
            filled_prompt = await getFilledPrompt(final_doc_obj)
            
            # 3. Add a small delay to prevent API flooding (429 errors)
            await asyncio.sleep(1)

            # 4. Send to Cerebras AI for content generation
            ai_response_text = await prompt_to_cerebras(filled_prompt)
            
            # 4. Parse AI response
            ai_content = parse_ai_json(ai_response_text)
            if not ai_content:
                raise HTTPException(status_code=500, detail="AI returned invalid content format")

            # 5. Trigger document generation
            doc_type = final_doc_obj.get("docType", "document").lower()
            
            if doc_type == "presentation":
                file_stream = create_ppt(ai_content)
                file_name = f"{ai_content.get('title', 'presentation')}.pptx"
                media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            else:
                file_stream = create_docx(ai_content)
                file_name = f"{ai_content.get('title', 'document')}.docx"
                media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

            return StreamingResponse(
                file_stream,
                media_type=media_type,
                headers={
                    "Content-Disposition": f"attachment; filename={file_name}"
                }
            )

        except Exception as e:
            print(f"ERROR in handle_generate_doc: {str(e)}")
            raise HTTPException(status_code=500, detail=str(e))

async def handle_process_file(file: UploadFile, topic_name: str):
    """
    Controller logic for processing uploaded files.
    """
    try:
        content = await file.read()
        filename = file.filename.lower()
        extracted_text = ""

        if filename.endswith(".docx"):
            doc = DocxDocument(BytesIO(content))
            extracted_text = "\n".join([para.text for para in doc.paragraphs])
        elif filename.endswith(".pptx"):
            prs = Presentation(BytesIO(content))
            text_runs = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_runs.append(shape.text)
            extracted_text = "\n".join(text_runs)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format")

        return {
            "message": "File processed successfully",
            "topicName": topic_name,
            "extractedContent": extracted_text,
            "filename": file.filename
        }
    except Exception as e:
        print(f"ERROR in handle_process_file: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))
