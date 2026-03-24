from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import uuid

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

def create_docx(response):
    
    print("i came to doc service")
    
    def style_run(run, bold=False, fSize=12, color=None):
        run.bold = bold
        run.font.size = Pt(fSize)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
            
    def style_para(para, align="left", space_after=0, line_spacing=1.5, first_indent=0):
        align_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        para.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.line_spacing = line_spacing
        para.paragraph_format.first_line_indent = Pt(first_indent)
    
    doc = Document()

    title = doc.add_paragraph()
    tRun = title.add_run(response["title"])
    style_para(title, "center", 15, 1.5, 0)
    style_run(tRun, True, 22, "FF0000")
    
    for section in response["sections"]:
        doc.add_paragraph()
        
        secHeading = doc.add_paragraph()
        secHeadRun = secHeading.add_run(section["sectionName"])
        style_para(secHeading, "center", 12, 1.5)
        style_run(secHeadRun, True, 18, "FF0000")
        
        pages = section["pages"]
        
        for i, page in enumerate(pages):
            pageHeading = doc.add_paragraph()
            pageRun = pageHeading.add_run(page["pageTitle"])
            style_para(pageHeading, "center", 6, 2)
            style_run(pageRun, True, 14, "0000FF")
            
            for page_txt in page["content"]:
                pagePara = doc.add_paragraph()
                pageParaRun = pagePara.add_run(page_txt)
                style_para(pagePara, "justify", 10, 1.5, 24)
                style_run(pageParaRun, False, 12)
                
        doc.add_page_break()

    file_name = f"{uuid.uuid4()}.docx"
    output_path = os.path.join(BASE_DIR, "generated_files", file_name)
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)

    return output_path